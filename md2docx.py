#!/usr/bin/env python3
"""md2docx - convert Markdown files to Word (.docx) documents.

Usage:
    python md2docx.py input.md                 # writes input.docx
    python md2docx.py input.md -o report.docx  # custom output name
    python md2docx.py a.md b.md c.md           # batch convert
    python md2docx.py notes/ -r                # convert a folder (recursive)

Supported Markdown: headings, paragraphs, bold, italic, strikethrough,
inline code, links, bullet/numbered/nested lists, task lists, blockquotes,
fenced code blocks, tables, horizontal rules, and local images.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

CODE_FONT = "Consolas"
MAX_IMAGE_WIDTH = Inches(6.0)


# --------------------------------------------------------------------------- #
# Low-level helpers
# --------------------------------------------------------------------------- #
def add_hyperlink(paragraph, url: str, text: str, bold=False, italic=False):
    """Append a clickable hyperlink run to a paragraph."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    run.append(rpr)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def shade(element_pr, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def shade_paragraph(paragraph, fill: str):
    shade(paragraph._p.get_or_add_pPr(), fill)


def shade_cell(cell, fill: str):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_left_border(paragraph, color="AAAAAA", size=18, space=8):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    borders.append(left)
    ppr.append(borders)


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    ppr.append(borders)


# --------------------------------------------------------------------------- #
# Converter
# --------------------------------------------------------------------------- #
class MarkdownToDocx:
    def __init__(self, base_dir: Path, template: Path | None = None):
        self.base_dir = base_dir
        self.doc = Document(str(template)) if template else Document()
        self.md = MarkdownIt("commonmark", {"html": False}).enable(
            ["table", "strikethrough"]
        )
        self._style_names = {s.name for s in self.doc.styles}

    # -- public ------------------------------------------------------------- #
    def convert(self, text: str) -> Document:
        tokens = self.md.parse(text)
        self._render_blocks(tokens, 0, len(tokens))
        return self.doc

    # -- style helpers ------------------------------------------------------ #
    def _style(self, name: str, fallback: str = "Normal") -> str:
        return name if name in self._style_names else fallback

    # -- block rendering ---------------------------------------------------- #
    def _find_close(self, tokens, start):
        """Index of the closing token matching the opening token at `start`."""
        depth = 0
        for i in range(start, len(tokens)):
            depth += tokens[i].nesting
            if depth == 0:
                return i
        return len(tokens) - 1

    def _render_blocks(self, tokens, start, end, ctx=None):
        """Render tokens[start:end]. ctx carries list / quote state."""
        ctx = ctx or {"quote": 0, "list_stack": []}
        i = start
        while i < end:
            tok = tokens[i]
            t = tok.type

            if t == "heading_open":
                level = int(tok.tag[1])
                inline = tokens[i + 1]
                p = self.doc.add_paragraph(style=self._style(f"Heading {min(level, 9)}"))
                self._render_inline(p, inline.children)
                i = self._find_close(tokens, i) + 1

            elif t == "paragraph_open":
                inline = tokens[i + 1]
                p = self._new_paragraph(ctx)
                self._render_inline(p, inline.children)
                i = self._find_close(tokens, i) + 1

            elif t in ("bullet_list_open", "ordered_list_open"):
                close = self._find_close(tokens, i)
                ordered = t == "ordered_list_open"
                ctx["list_stack"].append({"ordered": ordered})
                self._render_blocks(tokens, i + 1, close, ctx)
                ctx["list_stack"].pop()
                i = close + 1

            elif t == "list_item_open":
                close = self._find_close(tokens, i)
                self._render_list_item(tokens, i, close, ctx)
                i = close + 1

            elif t == "blockquote_open":
                close = self._find_close(tokens, i)
                ctx["quote"] += 1
                self._render_blocks(tokens, i + 1, close, ctx)
                ctx["quote"] -= 1
                i = close + 1

            elif t in ("fence", "code_block"):
                self._render_code_block(tok.content)
                i += 1

            elif t == "hr":
                add_horizontal_rule(self.doc)
                i += 1

            elif t == "table_open":
                close = self._find_close(tokens, i)
                self._render_table(tokens, i, close)
                i = close + 1

            else:
                i += 1

    def _new_paragraph(self, ctx, style=None):
        if style is None:
            style = "Quote" if ctx["quote"] and "Quote" in self._style_names else "Normal"
        p = self.doc.add_paragraph(style=self._style(style))
        if ctx["quote"]:
            p.paragraph_format.left_indent = Inches(0.4 * ctx["quote"])
            set_left_border(p)
        return p

    def _render_list_item(self, tokens, start, end, ctx):
        stack = ctx["list_stack"]
        depth = min(len(stack), 3)
        info = stack[-1]
        base = "List Number" if info["ordered"] else "List Bullet"
        style = base if depth == 1 else f"{base} {depth}"

        first_paragraph = True
        i = start + 1
        while i < end:
            tok = tokens[i]
            if tok.type == "paragraph_open" and first_paragraph:
                inline = tokens[i + 1]
                p = self.doc.add_paragraph(style=self._style(style))
                children = inline.children or []
                # Task list: "[ ] " / "[x] " prefix
                if children and children[0].type == "text":
                    m = re.match(r"^\[([ xX])\]\s+", children[0].content)
                    if m:
                        box = "\u2612 " if m.group(1).lower() == "x" else "\u2610 "
                        p.add_run(box)
                        children[0].content = children[0].content[m.end():]
                self._render_inline(p, children)
                first_paragraph = False
                i = self._find_close(tokens, i) + 1
            elif tok.type in ("bullet_list_open", "ordered_list_open"):
                close = self._find_close(tokens, i)
                stack.append({"ordered": tok.type == "ordered_list_open"})
                self._render_blocks(tokens, i + 1, close, ctx)
                stack.pop()
                i = close + 1
            else:
                # Other block content inside a list item (extra paragraphs, code...)
                close = self._find_close(tokens, i) if tok.nesting == 1 else i
                saved = ctx["list_stack"]
                self._render_blocks(tokens, i, close + 1, {"quote": ctx["quote"], "list_stack": []})
                ctx["list_stack"] = saved
                i = close + 1

    def _render_code_block(self, code: str):
        lines = code.rstrip("\n").split("\n")
        for line in lines:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.left_indent = Inches(0.2)
            shade_paragraph(p, "F2F2F2")
            run = p.add_run(line if line else " ")
            run.font.name = CODE_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            run.font.size = Pt(9.5)
        self.doc.add_paragraph()  # spacing after block

    def _render_table(self, tokens, start, end):
        rows = []  # list of (is_header, [cell_inline_tokens, alignment])
        i = start
        current = None
        header = False
        while i < end:
            t = tokens[i]
            if t.type == "thead_open":
                header = True
            elif t.type == "tbody_open":
                header = False
            elif t.type == "tr_open":
                current = []
            elif t.type in ("th_open", "td_open"):
                align = "left"
                style_attr = t.attrGet("style") or ""
                m = re.search(r"text-align:\s*(\w+)", style_attr)
                if m:
                    align = m.group(1)
                current.append((tokens[i + 1].children or [], align))
            elif t.type == "tr_close":
                rows.append((header, current))
            i += 1

        if not rows:
            return
        ncols = max(len(r) for _, r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = self._style("Table Grid")

        from docx.enum.text import WD_ALIGN_PARAGRAPH
        align_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        for r, (is_header, cells) in enumerate(rows):
            for c, (children, align) in enumerate(cells):
                cell = table.cell(r, c)
                p = cell.paragraphs[0]
                p.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.LEFT)
                self._render_inline(p, children, force_bold=is_header)
                if is_header:
                    shade_cell(cell, "D9E2F3")
        self.doc.add_paragraph()

    # -- inline rendering --------------------------------------------------- #
    def _render_inline(self, paragraph, children, force_bold=False):
        bold = italic = strike = False
        code = False
        link = None

        for tok in children or []:
            t = tok.type
            if t == "strong_open":
                bold = True
            elif t == "strong_close":
                bold = False
            elif t == "em_open":
                italic = True
            elif t == "em_close":
                italic = False
            elif t == "s_open":
                strike = True
            elif t == "s_close":
                strike = False
            elif t == "link_open":
                link = tok.attrGet("href")
            elif t == "link_close":
                link = None
            elif t in ("softbreak",):
                paragraph.add_run(" ")
            elif t == "hardbreak":
                paragraph.add_run().add_break(WD_BREAK.LINE)
            elif t == "code_inline":
                run = paragraph.add_run(tok.content)
                run.font.name = CODE_FONT
                run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
                run.font.size = Pt(10)
                run.bold = bold or force_bold or None
                run.italic = italic or None
                shade(run._element.get_or_add_rPr(), "EDEDED")
            elif t == "image":
                self._add_image(paragraph, tok)
            elif t == "text":
                if link:
                    add_hyperlink(paragraph, link, tok.content,
                                  bold=bold or force_bold, italic=italic)
                else:
                    run = paragraph.add_run(tok.content)
                    run.bold = (bold or force_bold) or None
                    run.italic = italic or None
                    run.font.strike = strike or None

    def _add_image(self, paragraph, tok):
        src = tok.attrGet("src") or ""
        alt = tok.content or src
        if re.match(r"^https?://", src):
            paragraph.add_run(f"[image: {alt}]").italic = True
            return
        path = (self.base_dir / src).expanduser()
        if not path.exists():
            paragraph.add_run(f"[missing image: {src}]").italic = True
            return
        run = paragraph.add_run()
        pic = run.add_picture(str(path))
        if pic.width > MAX_IMAGE_WIDTH:
            ratio = MAX_IMAGE_WIDTH / pic.width
            pic.width = MAX_IMAGE_WIDTH
            pic.height = int(pic.height * ratio)


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def convert_file(src: Path, dst: Path, template: Path | None = None) -> Path:
    text = src.read_text(encoding="utf-8")
    doc = MarkdownToDocx(src.parent, template).convert(text)
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return dst


def collect_inputs(paths, recursive):
    files = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            pattern = "**/*.md" if recursive else "*.md"
            files.extend(sorted(p.glob(pattern)))
        elif p.is_file():
            files.append(p)
        else:
            print(f"warning: {p} not found, skipping", file=sys.stderr)
    return files


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Convert Markdown to Word (.docx)")
    ap.add_argument("inputs", nargs="+", help="Markdown file(s) or folder(s)")
    ap.add_argument("-o", "--output", help="Output file (single input) or folder (multiple)")
    ap.add_argument("-r", "--recursive", action="store_true", help="Recurse into folders")
    ap.add_argument("-t", "--template", help="Optional .docx/.dotx template for styles")
    args = ap.parse_args(argv)

    files = collect_inputs(args.inputs, args.recursive)
    if not files:
        print("No Markdown files found.", file=sys.stderr)
        return 1

    template = Path(args.template) if args.template else None
    out = Path(args.output) if args.output else None
    single_file_out = out is not None and out.suffix.lower() == ".docx"
    if single_file_out and len(files) > 1:
        print("error: -o must be a folder when converting multiple files", file=sys.stderr)
        return 2

    failures = 0
    for f in files:
        if single_file_out:
            dst = out
        elif out:
            dst = out / f.with_suffix(".docx").name
        else:
            dst = f.with_suffix(".docx")
        try:
            convert_file(f, dst, template)
            print(f"{f} -> {dst}")
        except Exception as exc:  # keep going on batch runs
            failures += 1
            print(f"error converting {f}: {exc}", file=sys.stderr)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
